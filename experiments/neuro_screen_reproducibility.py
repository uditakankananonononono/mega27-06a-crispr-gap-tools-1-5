"""Massively parallel neuronal-depolarization CRISPR screen (Nat Commun 2026,
10.1038/s41467-026-75882-0, MOESM4/8/9): 881-gene primary screen, 390-gene
multi-condition follow-up, 1,139 sgRNAs across 4+ conditions.
Gap-1: does guide-level effect size transfer across assay conditions, and do
sgRNAs targeting the same gene agree?"""
import json
import re
from collections import defaultdict
import numpy as np
import openpyxl
from scipy.stats import spearmanr

wb9 = openpyxl.load_workbook('data/raw/neuro_screen/moesm9.xlsx', read_only=True)
ws = wb9['Supplementary Data 9_crispr_scr']
rows = list(ws.iter_rows(values_only=True))
hdr = [str(c) for c in rows[0]]
# group columns by condition prefix
conds = defaultdict(dict)
for i, h in enumerate(hdr):
    m = re.match(r'(.+)\.sgrnaresult\.(log2_fold_change|fdr)$', h)
    if m:
        conds[m.group(1)][m.group(2)] = i

guides = []
for r in rows[1:]:
    if r[0] is None:
        continue
    rec = {'sgrna': str(r[0]), 'gene': str(r[1])}
    for cond, cols in conds.items():
        try:
            rec[cond] = float(r[cols['log2_fold_change']])
        except (TypeError, ValueError):
            rec[cond] = None
    guides.append(rec)
assert len(guides) >= 1100, len(guides)

# pairwise condition transfer (complete cases)
cond_names = sorted(conds)
pairwise = {}
for i, a in enumerate(cond_names):
    for b in cond_names[i + 1:]:
        xs = [(g[a], g[b]) for g in guides if g[a] is not None and g[b] is not None]
        if len(xs) >= 50:
            pairwise[f'{a}_vs_{b}'] = {
                'n': len(xs),
                'spearman': float(spearmanr([x for x, _ in xs], [y for _, y in xs]).statistic),
            }

# within-gene sgRNA sign concordance on the primary condition
bygene = defaultdict(list)
for g in guides:
    if g.get('primary') is not None:
        bygene[g['gene']].append(g['primary'])
multi = {g: v for g, v in bygene.items() if len(v) >= 2}
concordant = sum(1 for v in multi.values()
                 if all(x > 0 for x in v) or all(x < 0 for x in v))
within_gene = {'n_genes_multi_sgRNA': len(multi),
               'n_sign_concordant': concordant,
               'fraction': concordant / len(multi) if multi else None,
               'mean_abs_spread': float(np.mean([max(v) - min(v) for v in multi.values()]))}

# gene-level multi-condition (MOESM8): phenotype-score correlations
wb8 = openpyxl.load_workbook('data/raw/neuro_screen/moesm8.xlsx', read_only=True)
ws = wb8['Supplementary Data 8_crispr_scr']
rows = list(ws.iter_rows(values_only=True))
hdr = [str(c) for c in rows[0]]
gconds = defaultdict(dict)
for i, h in enumerate(hdr):
    m = re.match(r'(.+)\.(phenotype_score|fdr)$', h)
    if m:
        gconds[m.group(1)][m.group(2)] = i
gene_rows = []
for r in rows[1:]:
    if r[0] is None:
        continue
    rec = {'gene': str(r[0])}
    for cond, cols in gconds.items():
        try:
            rec[cond] = float(r[cols['phenotype_score']])
        except (TypeError, ValueError, KeyError):
            rec[cond] = None
    gene_rows.append(rec)
gcond_names = sorted(gconds)
gpair = {}
for i, a in enumerate(gcond_names):
    for b in gcond_names[i + 1:]:
        xs = [(g[a], g[b]) for g in gene_rows if g[a] is not None and g[b] is not None]
        if len(xs) >= 50:
            gpair[f'{a}_vs_{b}'] = {'n': len(xs),
                                    'spearman': float(spearmanr([x for x, _ in xs], [y for _, y in xs]).statistic)}

out = {
    'n_sgrna': len(guides),
    'sgrna_conditions': cond_names,
    'sgrna_condition_transfer': pairwise,
    'within_gene_concordance': within_gene,
    'n_genes_genelevel': len(gene_rows),
    'genelevel_conditions': gcond_names,
    'genelevel_condition_transfer': gpair,
}
print(json.dumps(out, indent=1))
json.dump(out, open('results/neuro_screen_reproducibility.json', 'w'), indent=1)
