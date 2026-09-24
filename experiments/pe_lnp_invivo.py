"""Prime editing with LNPs in vivo/in vitro (Nat Nanotechnol 2026,
10.1038/s41565-026-02200-6, MOESM5): epegRNA vs HM-pegRNA dose-response
(Fig 1b), PE editor ladder at Pcsk9 (Fig 2a), editing kinetics (Fig 2c).
Gap-3: delivery/editor/dose axes of PE efficiency."""
import json
import numpy as np
import openpyxl

wb = openpyxl.load_workbook('data/raw/pe_lnp/moesm5.xlsx', read_only=True)
out = {}

# Fig 1b: dose-response, cols 1-3 HM-pegRNA, 4-6 epegRNA
ws = wb['Figure 1b']
doses = []
for r in list(ws.iter_rows(values_only=True))[1:]:
    if isinstance(r[0], (int, float)):
        hm = [float(x) for x in r[1:4] if isinstance(x, (int, float))]
        ep = [float(x) for x in r[4:7] if isinstance(x, (int, float))]
        if hm and ep:
            doses.append(dict(dose=float(r[0]), hm=round(float(np.mean(hm)), 2),
                              epeg=round(float(np.mean(ep)), 2),
                              ratio=round(float(np.mean(ep)) / max(float(np.mean(hm)), 1e-9), 2)))
out['dose_response'] = doses

# Fig 2a: editor ladder, cols 1-3 edit, 4-6 indel
ws = wb['Figure 2a']
editors = []
for r in list(ws.iter_rows(values_only=True))[1:]:
    if r[0] is not None:
        edit = [float(x) for x in r[1:4] if isinstance(x, (int, float))]
        ind = [float(x) for x in r[4:7] if isinstance(x, (int, float))]
        editors.append(dict(editor=str(r[0]).replace('\xa0', ' '),
                            edit=round(float(np.mean(edit)), 2) if edit else None,
                            indel=round(float(np.mean(ind)), 3) if ind else None))
out['editor_ladder'] = editors

# Fig 2c: kinetics rows=time hours, cols 1-3 PEmax, 4-6 PE6c
ws = wb['Figure 2c']
kin = []
for r in list(ws.iter_rows(values_only=True))[1:]:
    if isinstance(r[0], (int, float)):
        pm = [float(x) for x in r[1:4] if isinstance(x, (int, float))]
        p6 = [float(x) for x in r[4:7] if isinstance(x, (int, float))]
        kin.append(dict(hours=float(r[0]),
                        pemax=round(float(np.mean(pm)), 3) if pm else None,
                        pe6c=round(float(np.mean(p6)), 3) if p6 else None))
out['kinetics'] = kin
json.dump(out, open('results/pe_lnp_invivo.json', 'w'), indent=1)
print(json.dumps(out, indent=1))
