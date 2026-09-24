"""Compact high-fidelity SaCas9 (Nat Commun 2026, 10.1038/s41467-026-71626-2,
MOESM7): PAM screen (Fig 2a indel, Fig 2b base editing, ~70 DYRK1A targets,
3 reps) for WT vs PAM-broadened SaCas9-NNG, plus off-target read panels
(Supp Fig 4d) for WT/eWT/NNG/eNNG. PAM-variant efficacy + fidelity trade-off."""
import json
import numpy as np
import openpyxl
from collections import defaultdict
from scipy.stats import spearmanr

wb = openpyxl.load_workbook('data/raw/sacas9_hf/moesm7.xlsx', read_only=True)

def parse_screen(sn):
    recs = []
    for r in wb[sn].iter_rows(min_row=3, values_only=True):
        if r[1] is None or r[4] is None:
            continue
        try:
            reps = [float(x) for x in r[6:9] if isinstance(x, (int, float))]
            recs.append(dict(enzyme=str(r[1]).strip(), target=str(r[2]), pam=str(r[4]),
                             mean=float(np.mean(reps)) if reps else np.nan))
        except (TypeError, ValueError):
            continue
    return recs

out = {}
for sn, tag in [('Fig. 2a', 'nuclease_indel'), ('Fig. 2b', 'base_editing')]:
    recs = parse_screen(sn)
    by_enz = defaultdict(list)
    for r in recs:
        by_enz[r['enzyme']].append(r)
    er = {}
    for enz, rs in by_enz.items():
        pam_mean = defaultdict(list)
        for r in rs:
            if not np.isnan(r['mean']):
                pam_mean[r['pam']].append(r['mean'])
        top = sorted(((p, round(float(np.mean(v)), 2), len(v)) for p, v in pam_mean.items()),
                     key=lambda x: -x[1])
        er[enz] = dict(n_targets=len(rs), top_pams=top[:5],
                       active_frac_gt10pct=round(sum(1 for r in rs if r['mean'] > 10) / len(rs), 3))
    out[tag] = er

# off-target panel Supp Fig 4d: groups WT/eWT/NNG/eNNG x 3 reps
ws = wb['Supplementary Fig. 4d']
hdr = None
sites = []
for r in ws.iter_rows(values_only=True):
    if r[2] and str(r[2]).startswith('chr'):
        groups = {}
        for gi, g in enumerate(['WT', 'eWT', 'NNG', 'eNNG']):
            cols = [5 + gi * 4, 6 + gi * 4, 7 + gi * 4]
            vals = [r[c] for c in cols if c < len(r) and isinstance(r[c], (int, float))]
            groups[g] = vals
        sites.append(groups)
out['offtarget_panel'] = {'n_sites': len(sites)}
for g in ['WT', 'eWT', 'NNG', 'eNNG']:
    totals = [sum(s[g]) if s[g] else 0.0 for s in sites]
    out['offtarget_panel'][g] = dict(
        total_reads=int(sum(totals)),
        sites_with_ge10_reads=sum(1 for t in totals if t >= 10),
        median_reads=round(float(np.median(totals)), 1))
wt_tot = np.array([sum(s['WT']) if s['WT'] else 0.0 for s in sites])
ewt_tot = np.array([sum(s['eWT']) if s['eWT'] else 0.0 for s in sites])
nng_tot = np.array([sum(s['NNG']) if s['NNG'] else 0.0 for s in sites])
enng_tot = np.array([sum(s['eNNG']) if s['eNNG'] else 0.0 for s in sites])
out['offtarget_panel']['eWT_read_reduction_frac'] = round(1 - ewt_tot.sum() / max(wt_tot.sum(), 1), 3)
out['offtarget_panel']['eNNG_read_reduction_frac'] = round(1 - enng_tot.sum() / max(nng_tot.sum(), 1), 3)
out['offtarget_panel']['wt_vs_ewt_spearman'] = round(float(spearmanr(wt_tot, ewt_tot).statistic), 3)

json.dump(out, open('results/sacas9_variants.json', 'w'), indent=1)
print(json.dumps(out, indent=1)[:3000])
