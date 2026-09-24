"""ABE9-SpRY bystander/off-target panel (Sci Rep 2026,
10.1038/s41598-026-40642-z, MOESM2): ACEofBASE-nominated off-targets for 4
mouse loci, measured cumulative A-to-G editing (S11; target + OT1-5 x
ABE9-SpRY/ABE8e-SpRY x triplicate + controls), XMAS-TREE FACS (S12).
Gap-2/4: does PAM-flexible SpRY pay an off-target cost, and how much does
the ABE9 deaminase buy back vs ABE8e?"""
import json
import numpy as np
import openpyxl

wb = openpyxl.load_workbook('data/raw/abe9_spry/moesm2.xlsx', read_only=True)

# nomination counts per target (S7-S10 data rows below the header block)
nom = {}
for sn, locus in (('S7. ACEofBASE Tpc1-I486 OTs', 'Tpc1-I486'),
                  ('S8. ACEofBASE Tpc2-K188 OTs', 'Tpc2-K188'),
                  ('S9. ACEofBASE Tpc2-L249 OTs', 'Tpc2-L249'),
                  ('S10. ACEofBASE Trpm4-L903 OTs', 'Trpm4-L903')):
    ws = wb[sn]
    n = 0
    for r in ws.iter_rows(values_only=True, min_row=15):
        if r[0] is not None and str(r[0]).strip() in [str(i) for i in range(1, 20)] + ['X', 'Y']:
            n += 1
    nom[locus] = n

# measured panel
ws = wb['S11. Cumulative OT raw data']
meas = {}
ctrl = {}
for r in ws.iter_rows(values_only=True, min_row=2):
    ed, locus, site, rep, val = r[:5]
    if ed is None:
        continue
    try:
        v = float(val)
    except (TypeError, ValueError):
        continue
    if ed == 'control':
        ctrl.setdefault((locus, site), []).append(v)
    else:
        meas.setdefault((locus, site, ed), []).append(v)

loci = sorted({k[0] for k in meas})
editors = sorted({k[2] for k in meas})
per_locus = {}
for locus in loci:
    rec = {}
    ctrl_mean = {s: float(np.mean(v)) for (l, s), v in ctrl.items() if l == locus}
    for ed in editors:
        on = float(np.mean(meas[(locus, 'target', ed)]))
        ots = {}
        for i in range(1, 6):
            site = f'OT{i}'
            key = (locus, site, ed)
            if key in meas:
                vals = meas[key]
                ots[site] = {'mean': float(np.mean(vals)),
                             'rep_cv': float(np.std(vals) / np.mean(vals)) if np.mean(vals) > 0 else None,
                             'control': ctrl_mean.get(site)}
        above = [s for s, x in ots.items()
                 if x['control'] is not None and x['mean'] > 3 * x['control']]
        rec[ed] = {'on_target': on, 'ots': ots,
                   'max_ot': max(x['mean'] for x in ots.values()),
                   'n_ot_above_3x_control': len(above)}
    per_locus[locus] = rec

reduction = {}
for locus in loci:
    a9 = per_locus[locus].get('ABE9-SpRY')
    a8 = per_locus[locus].get('ABE8e-SpRY')
    if a9 and a8:
        reduction[locus] = {
            'on_target_fold_8e_over_9': a8['on_target'] / max(a9['on_target'], 1e-9),
            'max_ot_fold_8e_over_9': a8['max_ot'] / max(a9['max_ot'], 1e-9),
            'n_ot_above_3x_ctrl_8e': a8['n_ot_above_3x_control'],
            'n_ot_above_3x_ctrl_9': a9['n_ot_above_3x_control'],
        }

out = {
    'aceofbase_nominated_per_locus': nom,
    'n_loci': len(loci),
    'per_locus': per_locus,
    'reduction_8e_vs_9': reduction,
    'note': 'cumulative A-to-G editing (%) at ACEofBASE-nominated sites; control = untreated amplicon background.',
}
print(json.dumps(out, indent=1))
json.dump(out, open('results/abe9_spry_offtarget.json', 'w'), indent=1)
