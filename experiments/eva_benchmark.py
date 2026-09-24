"""EVA score cross-study evaluation (Nat Commun 2025,
10.1038/s41467-025-59947-0, MOESM6): per-study Spearman of the paper's
EVA activity score vs measured activity across 10 re-deposited studies.
The deposit counts once; re-deposited older studies are used only to
score the new model (not recounted)."""
import json
import openpyxl
import numpy as np
from scipy.stats import spearmanr

wb = openpyxl.load_workbook('data/raw/grna_disentangle/moesm6.xlsx', read_only=True)
out = {}
for sn in wb.sheetnames:
    ws = wb[sn]
    act, eva = [], []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        try:
            a, e = float(r[1]), float(r[2])
        except (TypeError, ValueError, IndexError):
            continue
        act.append(a); eva.append(e)
    if len(act) >= 10:
        rho = spearmanr(act, eva)
        out[sn] = dict(n=len(act), spearman=round(float(rho.statistic), 3),
                       pvalue=float(rho.pvalue))
json.dump(out, open('results/eva_benchmark.json', 'w'), indent=1)
print(json.dumps(out, indent=1))
