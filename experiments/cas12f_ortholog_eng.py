"""Comparative Cas12f orthologs (Nat Struct Mol Biol 2026,
10.1038/s41594-026-01788-6, MOESM9 Fig 3d + MOESM10 Fig 4l): guide-RNA
truncation tolerance (fold interference per gRNA deletion variant) and
interface-mutant sensitivity (normalized fold per point mutant)."""
import json
import numpy as np
import openpyxl

out = {}
wb = openpyxl.load_workbook('data/raw/cas12f_orthologs/moesm9.xlsx', read_only=True)
ws = wb['Sheet1']
guide_vars = []
for r in ws.iter_rows(min_row=4, values_only=True):
    if r[0] is None:
        continue
    vals = [float(x) for x in r[1:4] if isinstance(x, (int, float))]
    if vals:
        guide_vars.append(dict(variant=str(r[0]), mean_fold=round(float(np.mean(vals)), 1)))
out['guide_variants'] = guide_vars
# split into per-ortholog blocks at each 'gRNA WT' row, normalize within block
blocks = []
cur = None
for g in guide_vars:
    if g['variant'] == 'gRNA WT':
        cur = []
        blocks.append(cur)
    if cur is not None:
        cur.append(g)
out['guide_truncation_retention'] = {}
for bi, blk in enumerate(blocks):
    wtv = blk[0]['mean_fold']
    for g in blk[1:]:
        out['guide_truncation_retention'][f"block{bi+1}:{g['variant']}"] = round(g['mean_fold'] / wtv, 3)

wb = openpyxl.load_workbook('data/raw/cas12f_orthologs/moesm10.xlsx', read_only=True)
ws = wb['Sheet1']
muts = []
for r in ws.iter_rows(min_row=4, values_only=True):
    if r[0] is None:
        continue
    vals = [float(x) for x in r[1:4] if isinstance(x, (int, float))]
    if vals:
        muts.append(dict(mutant=str(r[0]), mean_fold=round(float(np.mean(vals)), 3)))
out['n_mutants'] = len(muts)
folds = [m['mean_fold'] for m in muts if m['mutant'] != 'WT']
out['mutant_fold_median'] = round(float(np.median(folds)), 3)
out['mutant_fold_min'] = round(min(folds), 3)
out['mutants_below_half'] = sum(1 for f in folds if f < 0.5)
out['most_sensitive'] = [m['mutant'] for m in sorted(muts, key=lambda x: x['mean_fold'])[:5] if m['mutant'] != 'WT']
json.dump(out, open('results/cas12f_ortholog_eng.json', 'w'), indent=1)
print(json.dumps(out, indent=1))
