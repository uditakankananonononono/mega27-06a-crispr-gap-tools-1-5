"""AlphaFold3 contact-model guided base editing (Nature 2026,
10.1038/s41586-026-10794-z, MOESM5/7): per-site measured off-target editing
(Fig 1f/g) joined to structural-model scores and change labels for ~2,000
candidate sites per panel across 4 guides (Fig 3a/3e).
Gap-2: do AF3-derived contact/structure differences separate measured
edited off-targets from silent ones?"""
import json
import numpy as np
import openpyxl

wb5 = openpyxl.load_workbook('data/raw/af3_base_editing/moesm5.xlsx', read_only=True)
wb7 = openpyxl.load_workbook('data/raw/af3_base_editing/moesm7.xlsx', read_only=True)

def norm_region(x):
    return str(x).rstrip('_')

# measured: max off_target_edit per (guide, region)
measured = {}
for sn in ('Fig. 1f', 'Fig. 1g'):
    ws = wb5[sn]
    for r in ws.iter_rows(values_only=True, min_row=2):
        if r[0] is None:
            continue
        key = (str(r[0]), norm_region(r[1]))
        try:
            v = float(r[4])
        except (TypeError, ValueError):
            continue
        measured[key] = max(measured.get(key, 0.0), v)

# structural panels
panels = {}
for sn in ('Fig. 3a', 'Fig. 3e'):
    ws = wb7[sn]
    recs = {}
    for r in ws.iter_rows(values_only=True, min_row=2):
        if r[0] is None:
            continue
        key = (str(r[0]), norm_region(r[1]))
        try:
            recs[key] = {'struc_ma_diff': float(r[2]), 'cp_ma_diff': float(r[3]),
                         'label': str(r[4])}
        except (TypeError, ValueError):
            continue
    panels[sn] = recs

def auroc(scores, labels):
    pos = [s for s, l in zip(scores, labels) if l]
    neg = [s for s, l in zip(scores, labels) if not l]
    if not pos or not neg:
        return None
    wins = sum(1 for p in pos for n in neg if p > n) + 0.5 * sum(1 for p in pos for n in neg if p == n)
    return wins / (len(pos) * len(neg))

out_panels = {}
for sn, recs in panels.items():
    joined = [(key, v) for key, v in recs.items() if key in measured]
    labels_dist = {}
    for v in recs.values():
        labels_dist[v['label']] = labels_dist.get(v['label'], 0) + 1
    # class separation on joined subset: edited = measured >= 0.5
    js = [(v['struc_ma_diff'], v['cp_ma_diff'], measured[k], v['label']) for k, v in joined]
    edited = [m >= 0.5 for _, _, m, _ in js]
    by_class = {}
    for lab in sorted({x[3] for x in js}):
        vals = [m for *_, m, l in [(x[0], x[1], x[2], x[3]) for x in js] if l == lab]
        by_class[lab] = {'n': len(vals), 'mean_measured_edit': float(np.mean(vals)),
                         'frac_edited_ge_0.5': float(np.mean([m >= 0.5 for m in vals]))}
    out_panels[sn] = {
        'n_sites': len(recs),
        'n_guides': len({k[0] for k in recs}),
        'label_distribution': labels_dist,
        'joined_measured': len(js),
        'auroc_struc_ma_diff': auroc([x[0] for x in js], edited),
        'auroc_cp_ma_diff': auroc([x[1] for x in js], edited),
        'class_measured': by_class,
    }

out = this = {
    'n_measured_regions': len(measured),
    'measured_edit_summary': {'max': float(max(measured.values())),
                              'frac_ge_0.5': float(np.mean([v >= 0.5 for v in measured.values()]))},
    'panels': out_panels,
}
print(json.dumps(out, indent=1))
json.dump(out, open('results/af3_contact_offtarget.json', 'w'), indent=1)
