"""Guide-reprogramming / tracrRNA-dependent off-targets (TDO) (Nat Commun 2026,
10.1038/s41467-026-74520-z): MOESM4 'Identificated TDO sites' (guide-independent
tracrRNA-driven GUIDE-seq signals) vs MOESM5 canonical GUIDE-seq off-target
sites (Tsai et al. SRR6012045-52). What fraction of canonical GUIDE-seq
off-target signal is guide-independent (TDO)?"""
import json
import numpy as np
import openpyxl
from collections import defaultdict

wb4 = openpyxl.load_workbook('data/raw/grna_reprogram/moesm4.xlsx', read_only=True)
tdo = []
for r in wb4['Identificated TDO sites'].iter_rows(min_row=2, values_only=True):
    if r[0] is None:
        continue
    try:
        tdo.append(dict(srr=str(r[0]), sample=str(r[1]), chrom=str(r[2]).replace('chr', ''),
                        pos=int(r[6]), reads=float(r[10]) if r[10] is not None else np.nan))
    except (TypeError, ValueError):
        continue

wb5 = openpyxl.load_workbook('data/raw/grna_reprogram/moesm5.xlsx', read_only=True)
canon = []
for r in wb5[wb5.sheetnames[0]].iter_rows(min_row=2, values_only=True):
    if r[0] is None:
        continue
    try:
        canon.append(dict(srr=str(r[1]), target=str(r[2]), chrom=str(r[3]).replace('chr', ''),
                          pos=int(r[9]), reads=float(r[11]) if r[11] is not None else np.nan))
    except (TypeError, ValueError):
        continue

out = dict(n_tdo=len(tdo), n_canonical=len(canon),
           tdo_samples=sorted({t['srr'] for t in tdo}),
           canon_samples=sorted({c['srr'] for c in canon}))
# per-sample overlap: TDO within 50bp of a canonical site in the same SRR sample
by_srr_c = defaultdict(list)
for c in canon:
    by_srr_c[c['srr']].append(c)
by_srr_t = defaultdict(list)
for t in tdo:
    by_srr_t[t['srr']].append(t)
per_sample = {}
tot_overlap = tot_tdo = 0
overlap_reads_share = {}
for srr, ts in sorted(by_srr_t.items()):
    cs = by_srr_c.get(srr, [])
    hit = 0
    tdo_reads_hit = tdo_reads_all = 0.0
    for t in ts:
        matched = any(c['chrom'] == t['chrom'] and abs(c['pos'] - t['pos']) <= 50 for c in cs)
        if matched:
            hit += 1
            if not np.isnan(t['reads']):
                tdo_reads_hit += t['reads']
        if not np.isnan(t['reads']):
            tdo_reads_all += t['reads']
    per_sample[srr] = dict(n_tdo=len(ts), n_canonical=len(cs), tdo_overlapping_canonical=hit,
                           frac_tdo_in_canonical=round(hit / len(ts), 3) if ts else None)
    tot_overlap += hit
    tot_tdo += len(ts)
    overlap_reads_share[srr] = round(tdo_reads_hit / tdo_reads_all, 3) if tdo_reads_all else None
out['per_sample'] = per_sample
out['total_tdo_overlapping_canonical'] = tot_overlap
out['overall_frac_tdo_in_canonical'] = round(tot_overlap / tot_tdo, 3) if tot_tdo else None
# reverse: fraction of canonical sites that carry a TDO signal
rev = 0
for c in canon:
    ts = by_srr_t.get(c['srr'], [])
    if any(t['chrom'] == c['chrom'] and abs(t['pos'] - c['pos']) <= 50 for t in ts):
        rev += 1
out['canonical_sites_with_tdo'] = rev
out['frac_canonical_with_tdo'] = round(rev / len(canon), 3) if canon else None
# reads concentration of TDO sites
reads = [t['reads'] for t in tdo if not np.isnan(t['reads'])]
out['tdo_reads_median'] = round(float(np.median(reads)), 1) if reads else None
out['tdo_reads_max'] = round(float(max(reads)), 1) if reads else None
json.dump(out, open('results/grna_reprogram_tdo.json', 'w'), indent=1)
print(json.dumps(out, indent=1))
