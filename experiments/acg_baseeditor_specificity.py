"""Simultaneous A/C/G base editors (Nat Commun 2026, 10.1038/s41467-026-76020-6,
MOESM10 Fig 3a + MOESM6): 11 genomic sites x 4 conditions (Untreated,
A&C-BE dUGI, smACG32, smACG34) x 7 conversion classes x 3 replicates.
Conversion-class purity per editor, and codon-reach expansion."""
import json
import numpy as np
import openpyxl
from collections import defaultdict

wb = openpyxl.load_workbook('data/raw/acg_baseeditor/moesm10.xlsx', read_only=True)
ws = wb['Fig. 3a']
rows = list(ws.iter_rows(values_only=True))

blocks = []
i = 0
while i < len(rows):
    r = rows[i]
    if r[0] and str(r[0]).startswith('(%)'):
        site = str(rows[i + 1][0]).strip()
        # class headers at this row, cols 1,4,7,...: find non-null beyond col0 on NEXT row? headers are on row i+1
        hdr_row = rows[i + 1]
        classes = []
        for c in range(1, len(hdr_row)):
            v = hdr_row[c]
            if v is not None and str(v).strip() and 'Rep' not in str(v):
                classes.append((c, str(v).replace('\n', ' ').strip()))
        # data rows follow: Untreated + 3 editors
        j = i + 3
        while j < len(rows) and rows[j][0] and not str(rows[j][0]).startswith('(%)'):
            editor = str(rows[j][0]).strip()
            if 'Rep' not in editor:
                means = {}
                for c0, cname in classes:
                    vals = [rows[j][c0 + k] for k in range(3)
                            if c0 + k < len(rows[j]) and isinstance(rows[j][c0 + k], (int, float))]
                    means[cname] = float(np.mean(vals)) if vals else 0.0
                blocks.append(dict(site=site, editor=editor, means=means))
            j += 1
        i = j
    else:
        i += 1

sites = sorted({b['site'] for b in blocks})
editors = sorted({b['editor'] for b in blocks} - {'Untreated'})
out = dict(n_sites=len(sites), sites=sites, editors=editors)
class_names = sorted({c for b in blocks for c in b['means']})
out['n_conversion_classes'] = len(class_names)

# per editor: mean per class across sites, and purity = intended class share
# merge spelling variants of the same editor
for b in blocks:
    if b['editor'] == 'A&CBE \u0394UGI':
        b['editor'] = 'A&C-BE \u0394UGI'
editors = sorted({b['editor'] for b in blocks} - {'Untreated'})

def ac_only(cls):
    terms = [t.strip() for t in cls.split(' and ')]
    return all(t.split('>')[0].strip() in ('A', 'C') for t in terms)

per_editor = {}
for enz in editors:
    recs = [b for b in blocks if b['editor'] == enz]
    cls_mean = {c: float(np.mean([b['means'].get(c, 0.0) for b in recs])) for c in class_names}
    total = sum(cls_mean.values())
    top_class = max(cls_mean, key=cls_mean.get)
    ac_sum = sum(v for c, v in cls_mean.items() if ac_only(c))
    g_sum = sum(v for c, v in cls_mean.items() if any(
        t.strip().startswith('G>') for t in c.split(' and ')))
    per_editor[enz] = dict(
        n_sites=len(recs),
        total_conversions_mean=round(total, 2),
        ac_only_share=round(ac_sum / total, 3) if total else None,
        g_editing_share=round(g_sum / total, 3) if total else None,
        top_class=top_class[:40],
        top_class_share=round(cls_mean[top_class] / total, 3) if total else None,
        class_means={c[:40]: round(v, 2) for c, v in
                     sorted(cls_mean.items(), key=lambda kv: -kv[1])})
out['per_editor'] = per_editor

# codon reach from MOESM6
wb6 = openpyxl.load_workbook('data/raw/acg_baseeditor/moesm6.xlsx', read_only=True)
reach = {}
for sn in wb6.sheetnames:
    editor = sn.replace('Conversions by ', '')
    n_aa = []
    for r in wb6[sn].iter_rows(min_row=3, values_only=True):
        if r[4]:
            n_aa.append(len([a for a in str(r[4]).split(',') if a.strip()]))
    if n_aa:
        reach[editor] = dict(n_codons=len(n_aa), mean_unique_aa=round(float(np.mean(n_aa)), 2),
                             max_unique_aa=int(max(n_aa)))
out['codon_reach'] = reach

json.dump(out, open('results/acg_baseeditor_specificity.json', 'w'), indent=1)
print(json.dumps(out, indent=1)[:2800])
