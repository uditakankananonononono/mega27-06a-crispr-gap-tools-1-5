"""Akcakaya 2018 (Nature 10.1038/s41586-018-0500-9, MOESM4/6): CIRCLE-seq
site lists for two targets in WT and KI conditions. Readcount decay by
mismatch count, bulge-carrying fraction, WT-vs-KI concordance."""
import json
import openpyxl
import numpy as np
from scipy.stats import spearmanr


def load(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[list(wb.sheetnames)[0]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        mmv = str(r[3])
        if mmv in ('NA', 'None', ''):
            continue
        rows.append(dict(coord=str(r[0]), wt=float(r[1] or 0), seq=str(r[2]),
                         mm=int(float(mmv)),
                         bulge=bool(str(r[4] or '').strip()
                                    and str(r[5]) not in ('NA', 'None')),
                         ki=float(r[6] or 0)))
    return rows

out = {}
for name, path in (('large', 'data/raw/akcakaya/moesm4.xlsx'),
                   ('small', 'data/raw/akcakaya/moesm6.xlsx')):
    rows = load(path)
    on = next((r for r in rows if r['mm'] == 0 and not r['bulge']), None)
    ref = on['wt'] if on else max(r['wt'] for r in rows)
    by_mm = {}
    for r in rows:
        by_mm.setdefault(r['mm'], []).append(r['wt'])
    decay = {k: dict(n=len(v), median_rel_wt=round(float(np.median(v)) / ref, 4))
             for k, v in sorted(by_mm.items())}
    nb = [r for r in rows if r['bulge']]
    conc = spearmanr([r['wt'] for r in rows], [r['ki'] for r in rows])
    out[name] = dict(n_sites=len(rows), n_bulge=len(nb),
                     bulge_frac=round(len(nb) / len(rows), 3),
                     decay_by_mm=decay,
                     wt_ki_spearman=round(float(conc.statistic), 3),
                     wt_ki_p=float(conc.pvalue))
json.dump(out, open('results/akcakaya_circleseq.json', 'w'), indent=1)
print(json.dumps(out, indent=1)[:2000])
