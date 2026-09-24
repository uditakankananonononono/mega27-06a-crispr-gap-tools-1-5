"""SURRO-seq (Liu et al. 2022, Nat Commun 10.1038/s41467-022-31543-6).
LibB (MOESM6 Data3.4): 7,250 candidate OTs x 110 therapeutic guides, IF% in
cells; CFD scoring + mismatch decay + AUROC. LibC (MOESM8 Data 5.3): 284
systematic 1-2 mismatch sites over 5 spacers; in-cell position gradient vs
CFD positional penalties, multiplicativity test for double mismatches."""
import json
import numpy as np
import openpyxl
from scipy.stats import mannwhitneyu, spearmanr
from sklearn.metrics import roc_auc_score
from crisprgap.cfd import cfd_score, _mm_pam_scores

_ACGT = set('ACGT')


def scorable(sp, off):
    return len(sp) == 20 and len(off) == 23 and set(sp) <= _ACGT and set(off) <= _ACGT


# ---- LibB ----
wb = openpyxl.load_workbook('data/raw/surroseq/moesm6.xlsx', read_only=True)
rows = list(wb['Data3.4 LibB  SURRO-seqOFF data'].iter_rows(min_row=4, values_only=True))
recs = []
for r in rows:
    if not (r[1] and r[2]):
        continue
    sp, off = str(r[1]).strip().upper(), str(r[2]).strip().upper()
    ifp = float(r[5]); sig = str(r[12]).strip() == 'Sig.'
    n_mm = sum(1 for a, b in zip(sp, off[:20]) if a != b)
    cfd = cfd_score(sp + 'NGG', off) if scorable(sp, off) else None
    recs.append(dict(sp=sp, n_mm=n_mm, ifp=ifp, sig=sig, cfd=cfd))

appl = [r for r in recs if r['cfd'] is not None]
y_sig = np.array([r['sig'] for r in appl])
cfd_arr = np.array([r['cfd'] for r in appl])
ifp_arr = np.array([r['ifp'] for r in appl])
out_b = dict(n_sites=len(recs), n_guides=len({r['sp'] for r in recs}),
             n_sig=int(y_sig.sum()),
             cfd_applicable=len(appl),
             auroc_cfd_sig=round(float(roc_auc_score(y_sig, cfd_arr)), 3),
             auroc_cfd_if01=round(float(roc_auc_score(ifp_arr >= 0.1, cfd_arr)), 3),
             auroc_cfd_if1=round(float(roc_auc_score(ifp_arr >= 1.0, cfd_arr)), 3),
             spearman_cfd_if=round(float(spearmanr(cfd_arr, ifp_arr)[0]), 3))
mm = {}
for r in recs:
    mm.setdefault(r['n_mm'], []).append(r)
out_b['mm_decay'] = {str(k): dict(n=len(v), sig_rate=round(sum(x['sig'] for x in v)/len(v), 3),
                                  median_if=float(np.median([x['ifp'] for x in v])))
                     for k, v in sorted(mm.items()) if k <= 6}

# ---- LibC ----
wb2 = openpyxl.load_workbook('data/raw/surroseq/moesm8.xlsx', read_only=True)
rows2 = list(wb2['Data 5.3 LibC SURRO-seq'].iter_rows(min_row=2, values_only=True))
single, double = [], []
for r in rows2:
    if not (r[1] and r[2]):
        continue
    sp, off = str(r[1]).strip().upper(), str(r[2]).strip().upper()
    mm_pos = [(i, sp[i], off[i]) for i in range(20) if sp[i] != off[i]]
    rec = dict(sp=sp, off=off, ifp=float(r[5]), mm=mm_pos,
               cfd=cfd_score(sp + 'NGG', off) if scorable(sp, off) else None)
    (single if len(mm_pos) == 1 else double if len(mm_pos) == 2 else []).append(rec)

mm_scores, _ = _mm_pam_scores()
# CFD mean positional penalty: average over all rX:dY at each position
pos_pen = {}
for key, v in mm_scores.items():
    p = int(key.split(',')[1])
    pos_pen.setdefault(p, []).append(v)
cfd_pos_mean = {p: float(np.mean(v)) for p, v in pos_pen.items()}

pos_if = {}
for r in single:
    pos_if.setdefault(r['mm'][0][0] + 1, []).append(r['ifp'])
common = sorted(set(pos_if) & set(cfd_pos_mean))
obs = [float(np.mean(pos_if[p])) for p in common]
pred = [cfd_pos_mean[p] for p in common]
out_c = dict(n_sites=len(single) + len(double), n_single=len(single), n_double=len(double),
             n_spacers=len({r['sp'] for r in single + double}),
             spearman_pos_if_vs_cfd_penalty=round(float(spearmanr(obs, pred)[0]), 3),
             spearman_cfd_if_all=round(float(spearmanr(
                 [r['cfd'] for r in single + double if r['cfd'] is not None],
                 [r['ifp'] for r in single + double if r['cfd'] is not None])[0]), 3))
# multiplicativity: double-mm observed vs product of constituent single-mm IF%
single_lookup = {(r['sp'], r['mm'][0]): r['ifp'] for r in single}
pairs = []
for r in double:
    a = single_lookup.get((r['sp'], r['mm'][0]))
    b = single_lookup.get((r['sp'], r['mm'][1]))
    if a is not None and b is not None:
        pairs.append((r['ifp'], a * b / 100.0))  # IF% product normalized
if len(pairs) > 5:
    out_c['double_mm_multiplicativity'] = dict(
        n=len(pairs), spearman_obs_vs_product=round(float(spearmanr([p[0] for p in pairs], [p[1] for p in pairs])[0]), 3))

json.dump(dict(libb=out_b, libc=out_c), open('results/surroseq.json', 'w'), indent=1)
print(json.dumps(dict(libb=out_b, libc=out_c), indent=1))
