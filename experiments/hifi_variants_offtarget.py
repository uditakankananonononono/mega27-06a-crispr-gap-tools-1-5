"""Increased-fidelity SpCas9 variants (Nat Commun 2023, 10.1038/s41467-023-41393-5).
MOESM8 Fig6b/7: amplicon indel% (3 replicates) for WT + HF variants across
GUIDE-seq-detected off-target amplicons. MOESM9 Summary: GUIDE-seq off-target
counts per variant per target site. Gap-4: how completely do HF variants
eliminate validated WT off-targets, and at what on-target cost?"""
import json
import numpy as np
import openpyxl

# --- amplicon indel table (forward-filled blocks) ---
wb = openpyxl.load_workbook('data/raw/hifi_rule/moesm8.xlsx', read_only=True)
ws = wb['Fig 6b, 7, Sup. Fig. 10']
rows = list(ws.iter_rows(min_row=2, values_only=True))
recs = []
cur_amp = cur_var = None
for r in rows:
    if r[0]:
        cur_amp = str(r[0]).strip()
    if r[1]:
        cur_var = str(r[1]).strip()
    if cur_amp and cur_var and r[3] is not None:
        try:
            recs.append((cur_amp, cur_var, float(r[3])))
        except (TypeError, ValueError):
            pass

from collections import defaultdict
means = defaultdict(list)
for a, v, x in recs:
    means[(a, v)].append(x)
mean_indel = {k: float(np.mean(v)) for k, v in means.items()}

variants = sorted({v for _, v in means})
# on-target amplicons end with ' ON'
on = {v: np.mean([x for (a, vv), x in mean_indel.items() if vv == v and a.endswith(' ON')])
      for v in variants if any(a.endswith(' ON') and vv == v for (a, vv) in mean_indel)}
off_by_var = defaultdict(list)
for (a, v), x in mean_indel.items():
    if not a.endswith(' ON') and v != 'dead SpCas9':
        off_by_var[v].append((a, x))

per_var = {}
for v, lst in off_by_var.items():
    xs = [x for _, x in lst]
    # retention vs WT for same amplicon
    ret = [x / mean_indel[(a, 'WT SpCas9')] for a, x in lst
           if (a, 'WT SpCas9') in mean_indel and mean_indel[(a, 'WT SpCas9')] > 0]
    per_var[v] = dict(n_offtarget_amplicons=len(lst),
                      frac_active_ge_0p1=round(float(np.mean([x >= 0.1 for x in xs])), 3),
                      median_retention_vs_wt=round(float(np.median(ret)), 4) if ret else None,
                      frac_eliminated_lt10pct=round(float(np.mean([r < 0.1 for r in ret])), 3) if ret else None,
                      on_target_mean=round(float(on.get(v, float('nan'))), 3))

# --- GUIDE-seq summary counts ---
wb9 = openpyxl.load_workbook('data/raw/hifi_rule/moesm9.xlsx', read_only=True)
rows9 = list(wb9['Summary'].iter_rows(min_row=2, values_only=True))
gs = defaultdict(dict)
cur_site = None
for r in rows9:
    if r[0]:
        cur_site = str(r[0]).strip()
    if r[1] and r[4] is not None:
        try:
            gs[cur_site][str(r[1]).strip()] = int(float(r[4]))
        except (TypeError, ValueError):
            pass
guide_seq_counts = {s: dict(v) for s, v in gs.items()}

out = dict(n_amplicon_variant_pairs=len(mean_indel), variants=variants,
           per_variant=per_var, guideseq_offtarget_counts=guide_seq_counts)
json.dump(out, open('results/hifi_variants_offtarget.json', 'w'), indent=1)
print(json.dumps({k: out[k] for k in ('n_amplicon_variant_pairs', 'variants')}, indent=1))
for v, d in per_var.items():
    print(v, d)
print(json.dumps(guide_seq_counts, indent=1)[:800])
