"""Sequence Display large-scale sequence-activity datasets (Nat Biotechnol 2026,
10.1038/s41587-026-03087-3, MOESM4 Fig 4E/5E): model activity predictions for
SlugCas9 PAM-context libraries - R2 vs training-set size (100/1000/16000,
3 seeds x 4 PAM contexts) and a model benchmark (Fig 4E).
Gap-1: how much data does sequence-activity prediction need before it beats
the mean, and which model class wins?"""
import json
import numpy as np
import openpyxl

wb = openpyxl.load_workbook('data/raw/seq_display/moesm4.xlsx', read_only=True)
rows = list(wb['Fig. 5E'].iter_rows(values_only=True))
PAMS = ['NNGA', 'NNGT', 'NNGC', 'NNGG']
SIZES = ['100', '1000', '16000']
r2 = {}  # r2[size][pam] = [3 seeds]
metric = None
for r in rows:
    if r[1] is not None and str(r[1]).strip() in ('R2', 'Pearson', 'Spearman'):
        metric = str(r[1]).strip()
        continue
    if metric == 'R2' and r[0] is not None and str(r[0]) in SIZES:
        size = str(r[0])
        vals = [float(x) for x in r[1:13]]
        r2[size] = {pam: vals[k * 3:(k + 1) * 3] for k, pam in enumerate(PAMS)}
assert len(r2) == 3, r2.keys()
size_mean = {s: float(np.mean([v for vv in r2[s].values() for v in vv])) for s in SIZES}
per_pam_16k = {pam: float(np.mean(r2['16000'][pam])) for pam in PAMS}
# Fig 4E: model benchmark on NNGA
mrows = list(wb['Fig. 4E'].iter_rows(values_only=True))
models = {}
for r in mrows:
    if r[0] is not None and str(r[0]).strip() == 'NNGA R2':
        continue
    hdr = [str(c) for c in mrows[1]]
models_row = None
for i, r in enumerate(mrows):
    if r[0] is not None and 'ESM-2' in str(r[1] if r[1] else ''):
        models_row = [str(c) for c in r]
        vals_row = mrows[i + 1]
        break
benchmark = {}
for k in range(len(models_row)):
    if models_row[k] is None or str(models_row[k]).strip() == '':
        continue
    name = str(models_row[k]).replace('\xa0', ' ').strip()
    try:
        benchmark[name] = float(vals_row[k])
    except (TypeError, ValueError):
        pass
out = {
    'r2_by_training_size_mean_over_pams_seeds': size_mean,
    'r2_at_16k_per_pam': per_pam_16k,
    'r2_at_100_fraction_negative': float(np.mean([v < 0 for vv in r2['100'].values() for v in vv])),
    'r2_at_1000_mean': float(np.mean([v for vv in r2['1000'].values() for v in vv])),
    'model_benchmark_NNGA_R2': benchmark,
}
print(json.dumps(out, indent=1))
json.dump(out, open('results/seq_display_learning.json', 'w'), indent=1)
