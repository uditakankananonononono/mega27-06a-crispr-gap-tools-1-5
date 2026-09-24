"""Directed evolution of pegRNA 3' RNA-stabilizing motifs (Nat Biotechnol 2026,
10.1038/s41587-026-03123-2, MOESM7 ST5): 19,000 pegRNA elements (7,000 parent
motifs, 4,000 combos, 4,000 hp-extended, 3,000 negatives, 1,000 double mutants)
with PE6c/PE7/PEmax editing and indel percentages, 4 reps each. Gap-3: which
3' motifs raise prime-editing efficiency, and does it transfer across editors?"""
import json
import numpy as np
import openpyxl
from collections import defaultdict
from scipy.stats import spearmanr

wb = openpyxl.load_workbook('data/raw/pe_motifs/moesm7.xlsx', read_only=True)
ws = wb[wb.sheetnames[0]]
hdr = next(ws.iter_rows(values_only=True))
idx = {h: i for i, h in enumerate(hdr) if h}

recs = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[0] is None:
        continue
    def f(col):
        v = r[idx[col]]
        return float(v) if isinstance(v, (int, float)) else np.nan
    recs.append(dict(
        vtype=str(r[idx['Motif_variant_type']]),
        motif=str(r[idx['Motif_name_full']])[:60],
        reads=f('PP_total_reads'),
        pe7=f('PE7_avg_Editing_pct'), pe7_indel=f('PE7_avg_Indels_pct'),
        pemax=f('PEmax_avg_Editing_pct'), pe6c=f('PE6c_avg_Editing_pct')))

out = dict(n_elements=len(recs))
usable = [r for r in recs if not np.isnan(r['reads']) and r['reads'] >= 100
          and not np.isnan(r['pe7'])]
out['n_usable_reads_ge100'] = len(usable)

by_type = defaultdict(list)
for r in usable:
    by_type[r['vtype']].append(r)
type_stats = {}
for vt, rs in sorted(by_type.items()):
    pe7 = np.array([r['pe7'] for rs_ in [rs] for r in rs_])
    type_stats[vt] = dict(
        n=len(rs),
        pe7_median=round(float(np.median(pe7)), 2),
        pe7_mean=round(float(np.mean(pe7)), 2),
        pe7_frac_gt10=round(float(np.mean(pe7 > 10)), 3),
        pemax_median=round(float(np.median([r['pemax'] for r in rs if not np.isnan(r['pemax'])])), 2),
        indel_median=round(float(np.median([r['pe7_indel'] for r in rs if not np.isnan(r['pe7_indel'])])), 2))
out['by_variant_type'] = type_stats

# top parent motifs (PARENT rows only, n >= 30 elements)
by_motif = defaultdict(list)
for r in usable:
    if r['vtype'] == 'PARENT':
        by_motif[r['motif']].append(r['pe7'])
motif_stats = sorted(
    ((m, len(v), round(float(np.median(v)), 2)) for m, v in by_motif.items() if len(v) >= 30),
    key=lambda x: -x[2])
out['top_parent_motifs'] = motif_stats[:8]
out['bottom_parent_motifs'] = motif_stats[-3:]

# cross-editor transfer
paired = [(r['pe7'], r['pemax']) for r in usable if not np.isnan(r['pemax'])]
out['pe7_vs_pemax_spearman'] = round(float(spearmanr([p[0] for p in paired], [p[1] for p in paired]).statistic), 3)
paired6 = [(r['pe7'], r['pe6c']) for r in usable if not np.isnan(r['pe6c'])]
out['pe7_vs_pe6c_spearman'] = round(float(spearmanr([p[0] for p in paired6], [p[1] for p in paired6]).statistic), 3)

json.dump(out, open('results/pe_motif_screen.json', 'w'), indent=1)
print(json.dumps(out, indent=1)[:2600])
